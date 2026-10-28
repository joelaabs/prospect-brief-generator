import pandas as pd, urllib.parse
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE as ILLEGAL
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

import argparse, json
ap = argparse.ArgumentParser(description="Build the Excel CRM from the master target list")
ap.add_argument("--master", required=True, help="master_target_list.csv from build_target_list.py")
ap.add_argument("--out", default="CRM.xlsx")
ap.add_argument("--prefill", help="optional JSON {company: {column: value}} of researched data")
args = ap.parse_args()
OUT = args.out
m = pd.read_csv(args.master)
m = m[m['tier'] != 'D - Low fit'].copy()
# same priority order as HubSpot batches: High fit first, A before B before C, biggest first
m['_f'] = m['glove_fit'].map({'High': 0, 'Medium': 1, 'Low': 2})
m = m.sort_values(['_f', 'tier', 'employees_at_fit_sites'], ascending=[True, True, False]).reset_index(drop=True)
hs = set()

LISTS = {
 'Prospect Status': ['Not started','Needs research','Ready to email','Emailed','Following up','Replied','Meeting set','Customer','Not now','Not a fit'],
 'Tier': ['A','B','C'],
 'Glove Fit': ['High','Medium','Low'],
 'Industry Group': ['Metalworking & Machinery','Construction','Other Manufacturing','Wholesale & Distribution','Logistics & Transportation',
                    'Plastics, Rubber & Chemicals','Automotive & Aerospace','Agriculture, Forestry & Landscaping','Services, Repair & Staffing',
                    'Healthcare','Food Processing','Public Sector & Municipal','Waste & Recycling','Energy, Utilities & Mining','Retail, Hospitality & Other'],
 'MI Region': ['Southeast (Detroit Metro)','West (Grand Rapids/Muskegon)','Southwest (Kalamazoo/Battle Creek)','Lansing / Mid-Michigan',
               'Saginaw / Bay / Thumb','Flint / Genesee','Upper Peninsula','Jackson / South Central','Northwest (Traverse City)','Northern Lower','Unknown'],
 'Contact Research': ['Not started','Search online','Call company','Found - ready','No contact found'],
 'Contact Role': ['Safety / EHS','Purchasing','Plant / Operations Mgr','Owner / Executive','HR','Maintenance','Office / Reception','Other'],
 'Email Version': ['A - Pain point','B - News/trigger'],
 'Next Step': ['Call','Email','Call + Email','Sample drop-off','Meeting','None'],
 'Reply Type': ['No reply yet','Positive','Not now','Not a fit','Referral'],
 'Yes/No': ['Yes','No'],
}

# columns: (header, width, group, list_name or None, kind)
C = [
 ('ID',6,'id',None,''),('Company',34,'id',None,''),('Prospect Status',16,'track','Prospect Status',''),
 ('Tier',6,'osha','Tier',''),('Glove Fit',9,'osha','Glove Fit',''),('Industry Group',26,'osha','Industry Group',''),
 ('Industry Detail',32,'osha',None,''),('MI Region',24,'osha','MI Region',''),('City',16,'osha',None,''),
 ('Street Address',26,'osha',None,''),('ZIP',7,'osha',None,''),('MI Employees',10,'osha',None,'int'),('MI Sites',7,'osha',None,'int'),
 ('OSHA Recordable Cases',11,'osha',None,'int'),('OSHA Injury Rate',9,'osha',None,'dec'),
 ('Website',26,'research',None,''),('Find Website',11,'research',None,'link'),('Find People (LinkedIn)',12,'research',None,'link'),
 ('Main Phone',15,'research',None,''),('General Email',24,'research',None,''),('Contact Research',15,'research','Contact Research',''),
 ('Contact 1 Name',18,'c1',None,''),('Contact 1 Title',20,'c1',None,''),('Contact 1 Role',18,'c1','Contact Role',''),
 ('Contact 1 Email',26,'c1',None,''),('Contact 1 Phone',15,'c1',None,''),
 ('Contact 2 Name',18,'c2',None,''),('Contact 2 Title',20,'c2',None,''),('Contact 2 Role',18,'c2','Contact Role',''),
 ('Contact 2 Email',26,'c2',None,''),('Contact 2 Phone',15,'c2',None,''),
 ('Email Version',15,'track','Email Version',''),('First Email Sent',12,'track',None,'date'),('Follow-up 1 Due',12,'track',None,'fdate'),
 ('Follow-up 2 Due',12,'track',None,'fdate'),('Next Step',13,'track','Next Step',''),('Next Step Date',12,'track',None,'date'),
 ('Last Contact Date',12,'track',None,'date'),('Reply Type',13,'track','Reply Type',''),('Meeting Booked',10,'track','Yes/No',''),
 ('Notes',50,'track',None,''),('In HubSpot',9,'track','Yes/No',''),
]
COL = {h: i+1 for i, (h, *_) in enumerate(C)}
GROUP_FILL = {'id':'1F3864','osha':'595959','research':'2E75B6','c1':'548235','c2':'548235','track':'C55A11'}
ARIAL = 'Arial'

wb = Workbook()
ws = wb.active; ws.title = 'Companies'
lists = wb.create_sheet('Lists')
for j, (name, vals) in enumerate(LISTS.items(), start=1):
    lists.cell(1, j, name).font = Font(name=ARIAL, bold=True)
    for i, v in enumerate(vals, start=2): lists.cell(i, j, v).font = Font(name=ARIAL)
    lists.column_dimensions[L(j)].width = max(14, max(len(v) for v in vals) + 2)
LREF = {name: f"=Lists!${L(j)}$2:${L(j)}${len(vals)+1}" for j, (name, vals) in enumerate(LISTS.items(), start=1)}

thin = Side(style='thin', color='BFBFBF')
for (h, w, g, ln, k), c in zip(C, range(1, len(C)+1)):
    cell = ws.cell(1, c, h)
    cell.font = Font(name=ARIAL, bold=True, color='FFFFFF', size=10)
    cell.fill = PatternFill('solid', fgColor=GROUP_FILL[g])
    cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.column_dimensions[L(c)].width = w
ws.row_dimensions[1].height = 42

PREFILL = json.load(open(args.prefill)) if args.prefill else {}

def g(q): return 'https://www.google.com/search?q=' + urllib.parse.quote_plus(q)
def li(q): return 'https://www.linkedin.com/search/results/people/?keywords=' + urllib.parse.quote(q)

datefmt = 'mm/dd/yyyy'
for i, r in m.iterrows():
    row = i + 2
    vals = {
     'ID': i+1, 'Company': r['company'], 'Prospect Status': 'Not started', 'Tier': r['tier'], 'Glove Fit': r['glove_fit'],
     'Industry Group': r['industry_group_broad'], 'Industry Detail': r['industry_group'], 'MI Region': r['mi_region'],
     'City': r['city'], 'Street Address': r['street_address'], 'ZIP': str(r['zip'])[:5] if pd.notna(r['zip']) else '',
     'MI Employees': int(r['mi_employees_osha']), 'MI Sites': int(r['mi_sites']), 'OSHA Recordable Cases': int(r['recordable_cases']),
     'OSHA Injury Rate': None if pd.isna(r['injury_rate_trir']) else float(r['injury_rate_trir']),
     'Contact Research': 'Not started', 'In HubSpot': 'Yes' if r['company'] in hs else 'No',
    }
    vals.update(PREFILL.get(r['company'], {}))
    for h, v in vals.items():
        if v is None or (isinstance(v, float) and pd.isna(v)): continue
        if isinstance(v, str): v = ILLEGAL.sub('', v).strip()
        ws.cell(row, COL[h], v)
    # links
    c = ws.cell(row, COL['Find Website'], 'Search'); c.hyperlink = g(f"{r['company']} {r['city']} Michigan"); c.font = Font(name=ARIAL, color='0563C1', underline='single', size=10)
    c = ws.cell(row, COL['Find People (LinkedIn)'], 'Search'); c.hyperlink = li(f"{r['company']} safety OR EHS OR purchasing"); c.font = Font(name=ARIAL, color='0563C1', underline='single', size=10)
    if vals.get('Website'):
        c = ws.cell(row, COL['Website']); c.hyperlink = vals['Website']; c.font = Font(name=ARIAL, color='0563C1', underline='single', size=10)
    # formulas
    s = f"{L(COL['First Email Sent'])}{row}"
    ws.cell(row, COL['Follow-up 1 Due'], f'=IF({s}="","",WORKDAY({s},3))')
    ws.cell(row, COL['Follow-up 2 Due'], f'=IF({s}="","",WORKDAY({s},7))')

N = len(m) + 1
MAXR = 7000  # room to add new companies
base = Font(name=ARIAL, size=10)
for row in ws.iter_rows(min_row=2, max_row=N):
    for c in row:
        if c.font is None or c.font.name != ARIAL or (c.font.color is None):
            if not c.hyperlink: c.font = base
for h, *_r in C:
    k = _r[3]
    col = L(COL[h])
    if k in ('date', 'fdate'):
        for r_ in range(2, MAXR+1): ws[f'{col}{r_}'].number_format = datefmt
    elif k == 'dec':
        for r_ in range(2, N+1): ws[f'{col}{r_}'].number_format = '0.0'
    elif k == 'int':
        for r_ in range(2, N+1): ws[f'{col}{r_}'].number_format = '#,##0'
    if h == 'Notes':
        for r_ in range(2, N+1): ws[f'{col}{r_}'].alignment = Alignment(wrap_text=False)

# dropdowns
for h, w, g_, ln, k in C:
    if ln:
        dv = DataValidation(type='list', formula1=LREF[ln], allow_blank=True, showErrorMessage=True,
                            errorTitle='Pick from the list', error=f'Choose a value from the {ln} list (edit options on the Lists tab).')
        dv.add(f'{L(COL[h])}2:{L(COL[h])}{MAXR}'); ws.add_data_validation(dv)
dvd = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True, showErrorMessage=True,
                     errorTitle='Date', error='Enter a date like 10/28/2026.')
for h in ['First Email Sent','Next Step Date','Last Contact Date']:
    dvd.add(f'{L(COL[h])}2:{L(COL[h])}{MAXR}')
ws.add_data_validation(dvd)

# status colors (badge-style)
SC = {'Not started':('EDEDED','404040'),'Needs research':('DDEBF7','1F4E79'),'Ready to email':('E4DFEC','5B2C6F'),
      'Emailed':('BDD7EE','1F3864'),'Following up':('FCE4D6','833C0B'),'Replied':('FFF2CC','7F6000'),
      'Meeting set':('E2EFDA','375623'),'Customer':('A9D08E','1E4620'),'Not now':('F8CBDA','7B1E44'),'Not a fit':('F4B6B6','8B0000')}
sc = L(COL['Prospect Status'])
for v, (bg, fg) in SC.items():
    ws.conditional_formatting.add(f'{sc}2:{sc}{MAXR}', CellIsRule(operator='equal', formula=[f'"{v}"'],
        fill=PatternFill('solid', fgColor=bg), font=Font(color=fg, bold=True)))
tc = L(COL['Tier'])
for v, bg in {'A':'A9D08E','B':'BDD7EE','C':'EDEDED'}.items():
    ws.conditional_formatting.add(f'{tc}2:{tc}{MAXR}', CellIsRule(operator='equal', formula=[f'"{v}"'], fill=PatternFill('solid', fgColor=bg), font=Font(bold=True)))
# follow-up due highlight (due today/overdue, no reply yet)
for h in ['Follow-up 1 Due','Follow-up 2 Due','Next Step Date']:
    col = L(COL[h]); rc = L(COL['Reply Type'])
    ws.conditional_formatting.add(f'{col}2:{col}{MAXR}', FormulaRule(
        formula=[f'AND({col}2<>"",{col}2<=TODAY(),OR({rc}2="",{rc}2="No reply yet"))'],
        fill=PatternFill('solid', fgColor='FFC000'), font=Font(bold=True)))

ws.freeze_panes = 'C2'
ws.auto_filter.ref = f'A1:{L(len(C))}{N}'

# ---------------- Dashboard ----------------
d = wb.create_sheet('Dashboard', 0)
S = "Companies!"
def rng(h): c_ = L(COL[h]); return f"{S}${c_}$2:${c_}${MAXR}"
H = Font(name=ARIAL, bold=True, size=14, color='1F3864'); B = Font(name=ARIAL, bold=True, size=10); N10 = Font(name=ARIAL, size=10)
hdr_fill = PatternFill('solid', fgColor='1F3864'); hdr_font = Font(name=ARIAL, bold=True, color='FFFFFF', size=10)
d['A1'] = 'Prospecting Dashboard'; d['A1'].font = H
d['A2'] = 'Updates automatically from the Companies tab. Do not type on this tab.'; d['A2'].font = Font(name=ARIAL, italic=True, size=9, color='7F7F7F')
kpis = [
 ('Companies on list', f'=COUNTA({rng("Company")})', '#,##0'),
 ('Emails sent (first email)', f'=COUNT({rng("First Email Sent")})', '#,##0'),
 ('Replies', f'=COUNTIFS({rng("Reply Type")},"<>",{rng("Reply Type")},"<>No reply yet")', '#,##0'),
 ('Reply rate', '=IF(B5=0,0,B6/B5)', '0.0%'),
 ('Meetings booked', f'=COUNTIF({rng("Meeting Booked")},"Yes")', '#,##0'),
 ('Customers', f'=COUNTIF({rng("Prospect Status")},"Customer")', '#,##0'),
 ('Follow-ups due today or overdue',
  f'=SUMPRODUCT(({rng("Follow-up 1 Due")}<>"")*(N(+{rng("Follow-up 1 Due")})<=TODAY())*(({rng("Reply Type")}="")+({rng("Reply Type")}="No reply yet")))', '#,##0'),
 ('Next steps due today or overdue', f'=COUNTIFS({rng("Next Step Date")},"<="&TODAY(),{rng("Prospect Status")},"<>Customer",{rng("Prospect Status")},"<>Not a fit")', '#,##0'),
]
d['A3'] = 'Key numbers'; d['A3'].font = B
for i, (lab, f, fmt) in enumerate(kpis):
    r = 4 + i
    d.cell(r, 1, lab).font = N10; c = d.cell(r, 2, f); c.font = Font(name=ARIAL, bold=True, size=11); c.number_format = fmt
# status table
r0 = 14
d.cell(r0, 1, 'Pipeline by Prospect Status').font = B
for j, t in enumerate(['Status','Companies','Tier A','Tier B','Tier C']):
    c = d.cell(r0+1, 1+j, t); c.font = hdr_font; c.fill = hdr_fill
for i, v in enumerate(LISTS['Prospect Status']):
    r = r0 + 2 + i
    d.cell(r, 1, v).font = N10
    d.cell(r, 2, f'=COUNTIF({rng("Prospect Status")},A{r})').font = N10
    for j, t in enumerate(['A','B','C']):
        d.cell(r, 3+j, f'=COUNTIFS({rng("Prospect Status")},$A{r},{rng("Tier")},"{t}")').font = N10
    bg, fg = SC[v]; d.cell(r,1).fill = PatternFill('solid', fgColor=bg); d.cell(r,1).font = Font(name=ARIAL, bold=True, color=fg, size=10)
# industry table
r1 = r0 + 2 + len(LISTS['Prospect Status']) + 2
d.cell(r1, 1, 'Results by Industry Group').font = B
heads = ['Industry Group','Companies','Emails sent','Replies','Reply rate','Meetings','Customers']
for j, t in enumerate(heads):
    c = d.cell(r1+1, 1+j, t); c.font = hdr_font; c.fill = hdr_fill
for i, v in enumerate(LISTS['Industry Group']):
    r = r1 + 2 + i
    d.cell(r, 1, v).font = N10
    d.cell(r, 2, f'=COUNTIF({rng("Industry Group")},A{r})')
    d.cell(r, 3, f'=COUNTIFS({rng("Industry Group")},A{r},{rng("First Email Sent")},"<>")')
    d.cell(r, 4, f'=COUNTIFS({rng("Industry Group")},A{r},{rng("Reply Type")},"<>",{rng("Reply Type")},"<>No reply yet")')
    d.cell(r, 5, f'=IF(C{r}=0,0,D{r}/C{r})').number_format = '0.0%'
    d.cell(r, 6, f'=COUNTIFS({rng("Industry Group")},A{r},{rng("Meeting Booked")},"Yes")')
    d.cell(r, 7, f'=COUNTIFS({rng("Industry Group")},A{r},{rng("Prospect Status")},"Customer")')
    for j in range(2, 8): d.cell(r, j).font = N10
# email version table
r2 = r1 + 2 + len(LISTS['Industry Group']) + 2
d.cell(r2, 1, 'A/B email test').font = B
for j, t in enumerate(['Email Version','Emails sent','Replies','Reply rate','Meetings']):
    c = d.cell(r2+1, 1+j, t); c.font = hdr_font; c.fill = hdr_fill
for i, v in enumerate(LISTS['Email Version']):
    r = r2 + 2 + i
    d.cell(r, 1, v).font = N10
    d.cell(r, 2, f'=COUNTIFS({rng("Email Version")},A{r},{rng("First Email Sent")},"<>")')
    d.cell(r, 3, f'=COUNTIFS({rng("Email Version")},A{r},{rng("Reply Type")},"<>",{rng("Reply Type")},"<>No reply yet")')
    d.cell(r, 4, f'=IF(B{r}=0,0,C{r}/B{r})').number_format = '0.0%'
    d.cell(r, 5, f'=COUNTIFS({rng("Email Version")},A{r},{rng("Meeting Booked")},"Yes")')
    for j in range(2, 6): d.cell(r, j).font = N10
for col, w in zip('ABCDEFG', [34,12,12,10,11,10,10]): d.column_dimensions[col].width = w

# ---------------- How to use ----------------
hw = wb.create_sheet('How to use', 1)
rows = [
 ('How to use this CRM', None),
 ('', None),
 ('1. Work in the Companies tab. One row = one company. Use the filter arrows in the header row to sort and filter (e.g. Tier = A, MI Region = West).', None),
 ('2. Column header colors tell you what to edit:', None),
 ('     Dark gray = OSHA data (filled in for you, leave as-is)', '595959'),
 ('     Blue = company research you fill in (website, main phone, general email, contact research status)', '2E75B6'),
 ('     Green = Contact 1 and Contact 2 (name, title, role, email, phone)', '548235'),
 ('     Orange = outreach tracking (status, email version, dates, reply, meeting, notes)', 'C55A11'),
 ('3. Dropdown cells: click the cell, then the small arrow. To change dropdown options, edit the Lists tab.', None),
 ('4. "Find Website" and "Find People (LinkedIn)" are clickable search links for each company.', None),
 ('5. Contact Research: Search online -> Call company -> Found - ready (or No contact found).', None),
 ('6. When you send the first email, type the date in First Email Sent (e.g. 10/28/2026). Follow-up 1 Due (+3 workdays) and Follow-up 2 Due (+7 workdays) fill in automatically and turn ORANGE when due and no reply yet.', None),
 ('7. Set Reply Type when someone answers, and update Prospect Status. Use Next Step + Next Step Date for calls, sample drop-offs and "Not now" recontact dates.', None),
 ('8. The Dashboard tab updates by itself: reply rate, meetings, follow-ups due, results by industry, and the A/B email test.', None),
 ('9. To add a new company, type it in the first empty row at the bottom. Dropdowns and formatting work down to row 7,000; copy the two Follow-up Due formulas down from the row above.', None),
 ('', None),
 ('Prospect Status meanings', None),
 ('Not started - on the list, nothing done | Needs research - missing contacts/emails | Ready to email - contact + email ready', None),
 ('Emailed - first email sent | Following up - in calls/follow-ups | Replied - they responded', None),
 ('Meeting set - call, visit or samples scheduled | Customer - ordering | Not now - try again later (set Next Step Date) | Not a fit - stop', None),
 ('', None),
 ('Data notes', None),
 ('Companies, tiers, employees, sites and injury numbers come from OSHA Injury Tracking Application Form 300A data: https://www.osha.gov/Establishment-Specific-Injury-and-Illness-Data', None),
 ('Tier is an ESTIMATE from Michigan employees at hands-on sites: A = 300+, B = 100-299, C = under 100 (starting thresholds; adjust as orders come in).', None),
 ('OSHA Injury Rate = recordable cases per 100 full-time workers per year. Injury numbers are for your background only - do not quote them in cold emails.', None),
]
for i, (t, color) in enumerate(rows, start=1):
    c = hw.cell(i, 1, t)
    c.font = Font(name=ARIAL, size=10, bold=(i == 1 or t in ('Prospect Status meanings','Data notes')), color=color or '000000')
    if i == 1: c.font = Font(name=ARIAL, size=14, bold=True, color='1F3864')
hw.column_dimensions['A'].width = 160

wb.move_sheet('Lists', offset=10)
wb.active = 0
wb.save(OUT)
print('rows', N-1)
